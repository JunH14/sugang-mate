from __future__ import annotations

import argparse
import csv
import difflib
import html
import json
import re
import struct
import sys
import zipfile
import zlib
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


FREESECT = 0xFFFFFFFF
ENDOFCHAIN = 0xFFFFFFFE


def clean_text(text: str) -> str:
    text = text.replace("\ufeff", " ")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def safe_stem(course_code: str, class_no: str, course_name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9가-힣_.-]+", "_", course_name).strip("_")
    stem = f"{course_code}_{class_no}"
    return f"{stem}_{name}" if name else stem


def resolve_project_path(project_dir: Path, value: str) -> Path | None:
    if not value:
        return None
    path = Path(value)
    if not path.is_absolute():
        path = project_dir / path
    return path


def extract_html_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"<script\b[^>]*>.*?</script>", " ", raw, flags=re.I | re.S)
    raw = re.sub(r"<style\b[^>]*>.*?</style>", " ", raw, flags=re.I | re.S)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(?:p|div|tr|table|li|h[1-6])>", "\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    return clean_text(html.unescape(raw))


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[page {index}]\n{text}")
    return clean_text("\n\n".join(pages))


def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document

        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return clean_text("\n".join(parts))
    except Exception:
        return extract_docx_text_from_zip(path)


def extract_docx_text_from_zip(path: Path) -> str:
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(path) as archive:
        xml_bytes = archive.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    parts: list[str] = []
    for paragraph in root.iter(f"{namespace}p"):
        texts = [node.text or "" for node in paragraph.iter(f"{namespace}t")]
        line = "".join(texts).strip()
        if line:
            parts.append(line)
    return clean_text("\n".join(parts))


@dataclass
class CfbEntry:
    name: str
    object_type: int
    start_sector: int
    size: int


class CfbReader:
    def __init__(self, data: bytes) -> None:
        self.data = data
        if data[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ValueError("not an OLE compound file")
        self.sector_size = 1 << struct.unpack_from("<H", data, 30)[0]
        self.mini_sector_size = 1 << struct.unpack_from("<H", data, 32)[0]
        self.first_dir_sector = struct.unpack_from("<I", data, 48)[0]
        self.mini_stream_cutoff = struct.unpack_from("<I", data, 56)[0]
        self.first_mini_fat_sector = struct.unpack_from("<I", data, 60)[0]
        self.num_mini_fat_sectors = struct.unpack_from("<I", data, 64)[0]
        self.first_difat_sector = struct.unpack_from("<I", data, 68)[0]
        self.num_difat_sectors = struct.unpack_from("<I", data, 72)[0]
        self.difat = self._read_difat()
        self.fat = self._read_fat()
        self.entries = self._read_directory()
        self.root = next((entry for entry in self.entries if entry.object_type == 5), None)
        self.mini_fat = self._read_mini_fat()
        self.mini_stream = self._read_regular_stream(self.root.start_sector, self.root.size) if self.root else b""

    def _sector(self, sector_id: int) -> bytes:
        start = (sector_id + 1) * self.sector_size
        return self.data[start : start + self.sector_size]

    def _read_difat(self) -> list[int]:
        difat = [
            value
            for value in struct.unpack_from("<109I", self.data, 76)
            if value not in (FREESECT, ENDOFCHAIN)
        ]
        sector = self.first_difat_sector
        for _ in range(self.num_difat_sectors):
            if sector in (FREESECT, ENDOFCHAIN):
                break
            chunk = self._sector(sector)
            entries_per_sector = self.sector_size // 4 - 1
            values = struct.unpack_from(f"<{entries_per_sector}I", chunk, 0)
            difat.extend(value for value in values if value not in (FREESECT, ENDOFCHAIN))
            sector = struct.unpack_from("<I", chunk, self.sector_size - 4)[0]
        return difat

    def _read_fat(self) -> list[int]:
        fat: list[int] = []
        entries_per_sector = self.sector_size // 4
        for sector in self.difat:
            if sector in (FREESECT, ENDOFCHAIN):
                continue
            fat.extend(struct.unpack_from(f"<{entries_per_sector}I", self._sector(sector), 0))
        return fat

    def _chain(self, start_sector: int, fat: list[int] | None = None) -> list[int]:
        if start_sector in (FREESECT, ENDOFCHAIN):
            return []
        table = fat or self.fat
        chain: list[int] = []
        seen: set[int] = set()
        sector = start_sector
        while sector not in (FREESECT, ENDOFCHAIN) and sector < len(table) and sector not in seen:
            seen.add(sector)
            chain.append(sector)
            sector = table[sector]
        return chain

    def _read_regular_stream(self, start_sector: int, size: int) -> bytes:
        chunks = [self._sector(sector) for sector in self._chain(start_sector)]
        return b"".join(chunks)[:size]

    def _read_mini_fat(self) -> list[int]:
        if self.first_mini_fat_sector in (FREESECT, ENDOFCHAIN) or not self.num_mini_fat_sectors:
            return []
        raw = b"".join(self._sector(sector) for sector in self._chain(self.first_mini_fat_sector))
        return list(struct.unpack_from(f"<{len(raw) // 4}I", raw, 0))

    def _read_mini_stream(self, start_sector: int, size: int) -> bytes:
        chunks: list[bytes] = []
        for sector in self._chain(start_sector, self.mini_fat):
            start = sector * self.mini_sector_size
            chunks.append(self.mini_stream[start : start + self.mini_sector_size])
        return b"".join(chunks)[:size]

    def _read_directory(self) -> list[CfbEntry]:
        raw = self._read_regular_stream(self.first_dir_sector, 10**9)
        entries: list[CfbEntry] = []
        for offset in range(0, len(raw), 128):
            entry = raw[offset : offset + 128]
            if len(entry) < 128:
                continue
            name_len = struct.unpack_from("<H", entry, 64)[0]
            if name_len < 2:
                continue
            name = entry[: name_len - 2].decode("utf-16le", errors="ignore")
            object_type = entry[66]
            start_sector = struct.unpack_from("<I", entry, 116)[0]
            size = struct.unpack_from("<Q", entry, 120)[0]
            entries.append(CfbEntry(name, object_type, start_sector, size))
        return entries

    def stream(self, name: str) -> bytes:
        entry = next((item for item in self.entries if item.name == name), None)
        if not entry:
            raise KeyError(name)
        if entry.size < self.mini_stream_cutoff and entry.object_type == 2:
            return self._read_mini_stream(entry.start_sector, entry.size)
        return self._read_regular_stream(entry.start_sector, entry.size)

    def stream_names(self) -> list[str]:
        return [entry.name for entry in self.entries if entry.object_type == 2]


def extract_hwp_text(path: Path) -> str:
    reader = CfbReader(path.read_bytes())
    header = reader.stream("FileHeader")
    compressed = bool(struct.unpack_from("<I", header, 36)[0] & 1)
    section_names = sorted(
        name for name in reader.stream_names() if re.fullmatch(r"Section\d+", name)
    )
    parts: list[str] = []
    for name in section_names:
        raw = reader.stream(name)
        if compressed:
            raw = zlib.decompress(raw, -15)
        parts.append(extract_hwp_records(raw))
    return clean_text("\n".join(parts))


def extract_hwp_records(data: bytes) -> str:
    offset = 0
    parts: list[str] = []
    while offset + 4 <= len(data):
        header = struct.unpack_from("<I", data, offset)[0]
        offset += 4
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if offset + 4 > len(data):
                break
            size = struct.unpack_from("<I", data, offset)[0]
            offset += 4
        payload = data[offset : offset + size]
        offset += size
        if tag_id == 67:
            text = payload.decode("utf-16le", errors="ignore")
            text = re.sub(r"[\u0000-\u001f]", " ", text)
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


def extract_attachment_text(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf_text(path), "pdf"
    if ext == ".docx":
        return extract_docx_text(path), "docx"
    if ext == ".hwp":
        return extract_hwp_text(path), "hwp"
    return "", f"unsupported:{ext}"


def course_header(row: dict[str, str]) -> str:
    lines = [
            f"과목명: {row.get('course_name', '')}",
            f"학수번호: {row.get('course_code', '')}",
            f"분반: {row.get('class_no', '')}",
            f"교수명: {row.get('professor', '')}",
            f"연도/학기: {row.get('year', '')} / {row.get('term', '')}",
            f"강의계획서 URL: {row.get('syllabus_url', '')}",
    ]
    if row.get("completion_type"):
        lines.insert(4, f"이수구분: {row['completion_type']}")
    if row.get("credit"):
        lines.insert(-1, f"학점: {row['credit']}")
    if row.get("class_hours"):
        lines.insert(-1, f"학점/시수: {row['class_hours']}")
    if row.get("schedule_summary"):
        lines.insert(-1, f"수업시간 및 강의실: {row['schedule_summary']}")
    return "\n".join(lines)


def schedule_entries_from_row(row: dict[str, str]) -> list[dict[str, object]]:
    raw = row.get("schedule_entries_json", "").strip()
    if not raw:
        return []
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return []
    return [item for item in parsed if isinstance(item, dict)] if isinstance(parsed, list) else []


def format_schedule_summary(entries: list[dict[str, object]]) -> str:
    return " / ".join(
        f"{entry.get('class_no', '')}분반 {entry.get('display', '')}".strip()
        for entry in entries
        if entry.get("display")
    )


def extract_completion_type(text: str, fallback: str = "") -> str:
    match = re.search(
        r"이수구분\s*[:：]?\s*(전공필수|전공선택|교양필수|교양선택|일반선택)",
        text,
        flags=re.I,
    )
    return match.group(1) if match else fallback.strip()


def build_course_text(project_dir: Path, row: dict[str, str]) -> tuple[str, str, int]:
    sections: list[str] = [course_header(row)]
    sources: list[str] = []

    attachment_path = resolve_project_path(project_dir, row.get("attachment_path", ""))
    if attachment_path and attachment_path.exists():
        try:
            attachment_text, source_type = extract_attachment_text(attachment_path)
            if attachment_text:
                sections.append(f"[첨부 강의계획서: {attachment_path.name}]\n{attachment_text}")
                sources.append(source_type)
        except Exception as error:
            sources.append(f"attachment_failed:{error}")

    html_path = resolve_project_path(project_dir, row.get("html_path", ""))
    if html_path and html_path.exists():
        try:
            html_text = extract_html_text(html_path)
            if html_text:
                sections.append(f"[강의계획안 HTML]\n{html_text}")
                sources.append("html")
        except Exception as error:
            sources.append(f"html_failed:{error}")

    combined = clean_text("\n\n".join(sections))
    return combined, ";".join(sources), len(combined)


def normalize_course_name(value: str) -> str:
    return "".join(character.casefold() for character in value if character.isalnum())


def normalize_for_similarity(text: str) -> str:
    # Ignore the generated metadata header because section number and URL differ.
    body = text.split("\n\n", 1)[1] if "\n\n" in text else text
    body = re.sub(r"https?://\S+", " ", body)
    body = re.sub(r"\s+", "", body)
    return body.casefold()


def text_similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    return difflib.SequenceMatcher(None, left, right, autojunk=False).ratio()


def deduplicate_candidates(
    candidates: list[dict[str, object]], threshold: float
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    groups: dict[str, list[dict[str, object]]] = {}
    for candidate in candidates:
        record = candidate["record"]
        assert isinstance(record, dict)
        key = normalize_course_name(str(record.get("course_name", "")))
        groups.setdefault(key, []).append(candidate)

    kept: list[dict[str, object]] = []
    duplicate_rows: list[dict[str, object]] = []
    for group in groups.values():
        ranked = sorted(
            group,
            key=lambda item: (
                "html" != str(item["record"].get("sources", "")),
                int(item["record"].get("char_count", 0)),
                -int(item["index"]),
            ),
            reverse=True,
        )
        group_kept: list[dict[str, object]] = []
        for candidate in ranked:
            record = candidate["record"]
            assert isinstance(record, dict)
            best_match: dict[str, object] | None = None
            best_score = 0.0
            for canonical in group_kept:
                score = text_similarity(
                    str(candidate["comparison_text"]), str(canonical["comparison_text"])
                )
                if score > best_score:
                    best_score = score
                    best_match = canonical

            if best_match is not None and best_score >= threshold:
                canonical_record = best_match["record"]
                assert isinstance(canonical_record, dict)
                duplicate_sections = canonical_record.setdefault("duplicate_sections", [])
                assert isinstance(duplicate_sections, list)
                duplicate_sections.append(
                    {
                        "course_code": record.get("course_code", ""),
                        "class_no": record.get("class_no", ""),
                        "professor": record.get("professor", ""),
                        "schedule_summary": record.get("schedule_summary", ""),
                        "schedule_entries": record.get("schedule_entries", []),
                        "syllabus_url": record.get("syllabus_url", ""),
                        "similarity": round(best_score, 6),
                    }
                )
                canonical_schedules = canonical_record.setdefault("schedule_entries", [])
                duplicate_schedules = record.get("schedule_entries", [])
                if isinstance(canonical_schedules, list) and isinstance(duplicate_schedules, list):
                    existing = {
                        json.dumps(item, ensure_ascii=False, sort_keys=True)
                        for item in canonical_schedules
                        if isinstance(item, dict)
                    }
                    for schedule in duplicate_schedules:
                        if not isinstance(schedule, dict):
                            continue
                        key = json.dumps(schedule, ensure_ascii=False, sort_keys=True)
                        if key not in existing:
                            canonical_schedules.append(schedule)
                            existing.add(key)
                    canonical_record["schedule_summary"] = format_schedule_summary(
                        [item for item in canonical_schedules if isinstance(item, dict)]
                    )
                duplicate_rows.append(
                    {
                        "course_name": record.get("course_name", ""),
                        "canonical_course_code": canonical_record.get("course_code", ""),
                        "canonical_class_no": canonical_record.get("class_no", ""),
                        "duplicate_course_code": record.get("course_code", ""),
                        "duplicate_class_no": record.get("class_no", ""),
                        "similarity": f"{best_score:.6f}",
                        "threshold": f"{threshold:.2f}",
                        "decision": "duplicate_removed",
                    }
                )
            else:
                group_kept.append(candidate)

        kept.extend(group_kept)

    kept.sort(key=lambda item: int(item["index"]))
    return kept, duplicate_rows


def write_jsonl(path: Path, rows: Iterable[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract syllabus text for RAG.")
    parser.add_argument("--project-dir", type=Path, default=Path("."), help="Project root directory.")
    parser.add_argument(
        "--duplicate-threshold",
        type=float,
        default=0.90,
        help="Remove same-name sections whose text similarity meets this threshold.",
    )
    args = parser.parse_args()

    project_dir = args.project_dir.resolve()
    courses_path = project_dir / "data" / "processed" / "courses.csv"
    text_dir = project_dir / "data" / "processed" / "text"
    text_dir.mkdir(parents=True, exist_ok=True)

    with courses_path.open(encoding="utf-8-sig", newline="") as handle:
        courses = list(csv.DictReader(handle))

    candidates: list[dict[str, object]] = []
    for index, row in enumerate(courses):
        stem = safe_stem(row.get("course_code", ""), row.get("class_no", ""), row.get("course_name", ""))
        text, sources, char_count = build_course_text(project_dir, row)
        # The current lecture-plan HTML is authoritative when it conflicts with
        # an older attached syllabus document.
        completion_type = row.get("completion_type", "").strip() or extract_completion_type(text)
        completion_type_source = "html" if row.get("completion_type", "").strip() else "document_text"
        schedule_entries = schedule_entries_from_row(row)
        text_path = text_dir / f"{stem}.txt"

        record = {
            "course_code": row.get("course_code", ""),
            "class_no": row.get("class_no", ""),
            "course_name": row.get("course_name", ""),
            "professor": row.get("professor", ""),
            "completion_type": completion_type,
            "completion_type_source": completion_type_source,
            "credit": row.get("credit", ""),
            "class_hours": row.get("class_hours", ""),
            "schedule_raw": row.get("schedule_raw", ""),
            "schedule_summary": row.get("schedule_summary", ""),
            "schedule_entries": schedule_entries,
            "syllabus_url": row.get("syllabus_url", ""),
            "text_path": str(text_path.relative_to(project_dir)),
            "sources": sources,
            "char_count": char_count,
            "text": text,
        }
        candidates.append(
            {
                "index": index,
                "row": row,
                "record": record,
                "comparison_text": normalize_for_similarity(text),
            }
        )
        print(f"{row.get('course_code')}-{row.get('class_no')} chars={char_count} sources={sources}")

    kept_candidates, duplicate_rows = deduplicate_candidates(
        candidates, threshold=args.duplicate_threshold
    )

    for old_text_path in text_dir.glob("*.txt"):
        old_text_path.unlink()

    records: list[dict[str, object]] = []
    report_rows: list[dict[str, object]] = []
    deduplicated_course_rows: list[dict[str, object]] = []
    for candidate in kept_candidates:
        record = candidate["record"]
        row = candidate["row"]
        assert isinstance(record, dict)
        assert isinstance(row, dict)
        text_path = project_dir / str(record["text_path"])
        text_path.write_text(str(record["text"]), encoding="utf-8")
        records.append(record)
        report_rows.append(
            {
                key: value
                for key, value in record.items()
                if key not in {"text", "duplicate_sections"}
            }
        )
        duplicate_sections = record.get("duplicate_sections", [])
        course_row = dict(row)
        course_row["schedule_summary"] = str(record.get("schedule_summary", ""))
        course_row["schedule_entries_json"] = json.dumps(
            record.get("schedule_entries", []), ensure_ascii=False
        )
        course_row["duplicate_class_nos"] = ",".join(
            str(item.get("class_no", ""))
            for item in duplicate_sections
            if isinstance(item, dict)
        )
        course_row["duplicate_count"] = (
            len(duplicate_sections) if isinstance(duplicate_sections, list) else 0
        )
        deduplicated_course_rows.append(course_row)

    write_jsonl(project_dir / "data" / "processed" / "syllabus_texts.jsonl", records)

    report_path = project_dir / "data" / "processed" / "text_extraction_report.csv"
    with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
        fieldnames = [
            "course_code",
            "class_no",
            "course_name",
            "professor",
            "completion_type",
            "completion_type_source",
            "credit",
            "class_hours",
            "schedule_raw",
            "schedule_summary",
            "schedule_entries",
            "syllabus_url",
            "text_path",
            "sources",
            "char_count",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(report_rows)

    duplicate_report_path = project_dir / "data" / "processed" / "duplicate_report.csv"
    duplicate_fieldnames = [
        "course_name",
        "canonical_course_code",
        "canonical_class_no",
        "duplicate_course_code",
        "duplicate_class_no",
        "similarity",
        "threshold",
        "decision",
    ]
    with duplicate_report_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=duplicate_fieldnames)
        writer.writeheader()
        writer.writerows(duplicate_rows)

    deduplicated_courses_path = project_dir / "data" / "processed" / "courses_deduplicated.csv"
    course_fieldnames = (
        list(courses[0].keys()) + ["duplicate_class_nos", "duplicate_count"]
        if courses
        else []
    )
    with deduplicated_courses_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=course_fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(deduplicated_course_rows)

    dedup_manifest = {
        "similarity_method": "difflib.SequenceMatcher",
        "same_name_only": True,
        "threshold": args.duplicate_threshold,
        "input_course_count": len(candidates),
        "output_course_count": len(records),
        "removed_duplicate_count": len(duplicate_rows),
    }
    (project_dir / "data" / "processed" / "dedup_manifest.json").write_text(
        json.dumps(dedup_manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(
        f"Done. input={len(candidates)} unique={len(records)} "
        f"duplicates_removed={len(duplicate_rows)} threshold={args.duplicate_threshold:.2f}"
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
